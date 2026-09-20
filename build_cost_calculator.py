#!/usr/bin/env python3
"""Build the AI Cost Calculator — a working .xlsx, not a screenshot of one.

    python build_cost_calculator.py

Three sheets:
    Calculator  — you type your usage, it computes monthly cost per model
    Rates       — the price table, EDITABLE, so the product does not rot when
                  vendors change pricing (they did twice in 2026)
    Read Me     — how to use it, and what the numbers do and do not include

Every result is a live formula referencing Rates, not a baked-in number. That is the whole
point: when Anthropic or OpenAI move a price, the buyer edits one cell and the sheet is
correct again. A calculator with hardcoded totals is a screenshot with extra steps.

Uploads to Google Sheets without modification.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUTDIR = r'C:\Projects\UserGuide\outputs'
OUT = os.path.join(OUTDIR, 'AI_Cost_Calculator.xlsx')

ORANGE = 'C4611F'
DARK = '1A1A24'
BAND = 'F4F0EA'
INPUT_BG = 'FFF4D6'          # yellow = "you type here"

H1 = Font(name='Calibri', size=16, bold=True, color=DARK)
H2 = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
BOLD = Font(name='Calibri', size=11, bold=True, color=DARK)
BODY = Font(name='Calibri', size=11, color=DARK)
MUTED = Font(name='Calibri', size=10, color='6B6660')

FILL_HEAD = PatternFill('solid', fgColor=ORANGE)
FILL_BAND = PatternFill('solid', fgColor=BAND)
FILL_INPUT = PatternFill('solid', fgColor=INPUT_BG)

THIN = Side(style='thin', color='D8D3CC')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MONEY = '"$"#,##0.00'
PCT = '0%'
NUM = '#,##0.0'

# model, id, input $/M, output $/M, cache-read $/M
# Verified July 2026. Cache read is ~10% of the input rate.
MODELS = [
    ('Claude Haiku 4.5',  'claude-haiku-4-5-20251001', 1.00,  5.00,  0.10),
    ('Claude Sonnet 4.6', 'claude-sonnet-4-6',         3.00, 15.00,  0.30),
    ('Claude Opus 4.8',   'claude-opus-4-8',           5.00, 25.00,  0.50),
]


def style_header(ws, row, last_col):
    for c in range(1, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = H2
        cell.fill = FILL_HEAD
        cell.border = BOX
        cell.alignment = Alignment(vertical='center')
    ws.row_dimensions[row].height = 20


def build_rates(wb):
    ws = wb.create_sheet('Rates')
    ws['A1'] = 'Model rates — edit these when pricing changes'
    ws['A1'].font = H1
    ws['A2'] = ('Per million tokens. These are the ONLY numbers you should need to '
                'update; the Calculator reads from here.')
    ws['A2'].font = MUTED

    heads = ['Model', 'Model ID', 'Input $/M', 'Output $/M', 'Cache read $/M']
    for i, h in enumerate(heads, start=1):
        ws.cell(row=4, column=i, value=h)
    style_header(ws, 4, len(heads))

    for r, (name, mid, ci, co, cc) in enumerate(MODELS, start=5):
        ws.cell(row=r, column=1, value=name).font = BOLD
        ws.cell(row=r, column=2, value=mid).font = BODY
        for col, val in ((3, ci), (4, co), (5, cc)):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = BODY
            cell.number_format = MONEY
            cell.fill = FILL_INPUT
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BOX

    ws['A9'] = 'Discounts'
    ws['A9'].font = BOLD
    ws['A10'] = 'Batch API discount'
    ws['B10'] = 0.50
    ws['B10'].number_format = PCT
    ws['B10'].fill = FILL_INPUT
    ws['C10'] = 'Async, up to 24h turnaround'
    ws['C10'].font = MUTED
    for r in (10,):
        for c in range(1, 4):
            ws.cell(row=r, column=c).border = BOX

    for col, w in zip('ABCDE', (22, 30, 14, 14, 16)):
        ws.column_dimensions[col].width = w
    return ws


def build_calculator(wb):
    ws = wb.create_sheet('Calculator', 0)
    ws['A1'] = 'AI Cost Calculator'
    ws['A1'].font = H1
    ws['A2'] = 'Type your monthly usage in the yellow cells. Everything else calculates.'
    ws['A2'].font = MUTED

    ws['A4'] = 'YOUR USAGE'
    ws['A4'].font = BOLD

    inputs = [
        ('Input tokens per month (millions)', 50, NUM,
         'Everything you send: prompts, files, conversation history'),
        ('Output tokens per month (millions)', 10, NUM,
         'Everything the model writes back'),
        ('Share of input served from cache', 0.0, PCT,
         'Cache hits cost ~10% of the input rate'),
        ('Share of traffic sent via Batch API', 0.0, PCT,
         'Async work only — not interactive requests'),
    ]
    for i, (label, default, fmt, note) in enumerate(inputs):
        r = 5 + i
        ws.cell(row=r, column=1, value=label).font = BODY
        cell = ws.cell(row=r, column=2, value=default)
        cell.font = BOLD
        cell.number_format = fmt
        cell.fill = FILL_INPUT
        cell.border = BOX
        ws.cell(row=r, column=3, value=note).font = MUTED

    ws['A11'] = 'MONTHLY COST BY MODEL'
    ws['A11'].font = BOLD

    heads = ['Model', 'Input cost', 'Output cost', 'Monthly total', 'Per day']
    for i, h in enumerate(heads, start=1):
        ws.cell(row=12, column=i, value=h)
    style_header(ws, 12, len(heads))

    for i, (name, _mid, _ci, _co, _cc) in enumerate(MODELS):
        r = 13 + i
        rates_row = 5 + i
        ws.cell(row=r, column=1, value=name).font = BOLD

        # Input: the non-cached share at full price, the cached share at the cache rate.
        ws.cell(row=r, column=2, value=(
            '=($B$5*(1-$B$7)*Rates!C{rr} + $B$5*$B$7*Rates!E{rr})'
            '*(1-$B$8*Rates!$B$10)').format(rr=rates_row))
        ws.cell(row=r, column=3, value=(
            '=$B$6*Rates!D{rr}*(1-$B$8*Rates!$B$10)').format(rr=rates_row))
        ws.cell(row=r, column=4, value='=B{r}+C{r}'.format(r=r))
        ws.cell(row=r, column=5, value='=D{r}/30'.format(r=r))

        for c in range(2, 6):
            cell = ws.cell(row=r, column=c)
            cell.number_format = MONEY
            cell.font = BOLD if c == 4 else BODY
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BOX
            if i % 2 == 0:
                ws.cell(row=r, column=c).fill = FILL_BAND

    last = 12 + len(MODELS)
    ws.cell(row=last + 2, column=1,
            value='Cheapest option').font = BOLD
    ws.cell(row=last + 2, column=2,
            value='=INDEX(A13:A{l},MATCH(MIN(D13:D{l}),D13:D{l},0))'.format(l=last)).font = BODY
    ws.cell(row=last + 3, column=1,
            value='Saving vs most expensive').font = BOLD
    sav = ws.cell(row=last + 3, column=2,
                  value='=MAX(D13:D{l})-MIN(D13:D{l})'.format(l=last))
    sav.number_format = MONEY
    sav.font = BODY

    ws.cell(row=last + 5, column=1, value=(
        'Routing mechanical work to the cheapest model is usually a bigger saving than '
        'any prompt optimisation.')).font = MUTED

    for col, w in zip('ABCDE', (36, 18, 16, 16, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A13'
    return ws


def build_readme(wb):
    ws = wb.create_sheet('Read Me')
    ws['A1'] = 'How to use this'
    ws['A1'].font = H1

    lines = [
        ('', ''),
        ('1.', 'Open the Calculator sheet. Type into the yellow cells only.'),
        ('2.', 'Read your monthly cost for each model in the table below them.'),
        ('3.', 'When a vendor changes pricing, edit the Rates sheet — nothing else.'),
        ('', ''),
        ('WHAT COUNTS AS A TOKEN', ''),
        ('', 'Roughly 4 characters of English, or about 0.75 words.'),
        ('', '1 million tokens is around 750,000 words.'),
        ('', 'Every turn resends the conversation, so long chats spend input tokens fast.'),
        ('', ''),
        ('WHAT THIS INCLUDES', ''),
        ('', 'Input tokens, output tokens, prompt-cache reads, and the Batch discount.'),
        ('', ''),
        ('WHAT IT DOES NOT INCLUDE', ''),
        ('', 'Cache WRITES, which cost more than a normal input token.'),
        ('', 'Subscription plans (Claude Pro/Max) — those are flat fees, not metered.'),
        ('', 'Failed or retried runs, which still bill for what they produced.'),
        ('', 'Tax.'),
        ('', ''),
        ('ACCURACY', ''),
        ('', 'Rates verified July 2026. Vendors moved pricing twice that year —'),
        ('', 'check the current rate card before making a large commitment.'),
        ('', ''),
        ('', 'Unofficial and independent. Not affiliated with, endorsed by, or'),
        ('', 'sponsored by Anthropic, OpenAI, GitHub or Microsoft.'),
        ('', 'etsy.com/shop/FranksMarketDesigns'),
    ]
    for i, (a, b) in enumerate(lines, start=3):
        ca = ws.cell(row=i, column=1, value=a)
        cb = ws.cell(row=i, column=2, value=b)
        ca.font = BOLD if a and not a[0].isdigit() else BODY
        cb.font = BODY
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 82
    return ws


# SUPERSEDED 2026-09-17 — page 1 of the shipped preview PDF is now rendered by
# covers/ and spliced in by rebuild_covers.py. This function is retained so a
# from-source rebuild still produces a complete document (page 2 is untouched);
# run rebuild_covers.py afterwards.
def build_preview_pdf():
    """A 2-page visual preview of the spreadsheet.

    The Etsy mockup pipeline composites real PDF pages, and the deliverable here is an
    .xlsx. This renders what the buyer will see so the listing images show the actual
    product rather than a generic icon. It is a preview, not the deliverable.
    """
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas as rl_canvas

    W, H = LETTER
    MX = 40
    path = os.path.join(OUTDIR, 'AI_Cost_Calculator_Preview.pdf')
    c = rl_canvas.Canvas(path, pagesize=LETTER)
    c.setTitle('AI Cost Calculator — preview')

    ink = HexColor('#1A1A24')
    org = HexColor('#C4611F')
    grey = HexColor('#6B6660')
    band = HexColor('#F4F0EA')
    yell = HexColor('#FFF4D6')

    def sheet(title, sub, heads, rows, input_rows=()):
        c.setFillColor(HexColor('#FFFFFF'))
        c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(org)
        c.rect(0, H - 10, W, 10, fill=1, stroke=0)
        c.setFillColor(ink)
        c.setFont('Helvetica-Bold', 24)
        c.drawString(MX, H - 60, title)
        c.setFillColor(grey)
        c.setFont('Helvetica', 11)
        c.drawString(MX, H - 78, sub)

        y = H - 120
        colw = (W - 2 * MX) / len(heads)
        c.setFillColor(org)
        c.rect(MX, y - 20, W - 2 * MX, 22, fill=1, stroke=0)
        c.setFillColor(HexColor('#FFFFFF'))
        c.setFont('Helvetica-Bold', 10)
        for i, h in enumerate(heads):
            c.drawString(MX + 8 + i * colw, y - 14, h)
        y -= 20

        for ri, row in enumerate(rows):
            rh = 22
            c.setFillColor(yell if ri in input_rows else (band if ri % 2 == 0 else HexColor('#FFFFFF')))
            c.rect(MX, y - rh, W - 2 * MX, rh, fill=1, stroke=0)
            for i, cell in enumerate(row):
                c.setFillColor(ink if i == 0 else grey)
                c.setFont('Helvetica-Bold' if i == 0 else 'Helvetica', 10)
                c.drawString(MX + 8 + i * colw, y - 15, str(cell))
            y -= rh
        return y

    y = sheet('AI Cost Calculator', 'Type in the yellow cells — everything else is a live formula',
              ['Your usage', 'Value', 'Note'],
              [('Input tokens / month (M)', '50.0', 'Prompts, files, history'),
               ('Output tokens / month (M)', '10.0', 'What the model writes back'),
               ('Cache hit share', '0%', 'Hits cost ~10% of input'),
               ('Batch API share', '0%', 'Async work only')],
              input_rows=(0, 1, 2, 3))
    y -= 30
    c.setFillColor(ink)
    c.setFont('Helvetica-Bold', 13)
    c.drawString(MX, y, 'MONTHLY COST BY MODEL')
    y -= 24
    colw = (W - 2 * MX) / 5
    c.setFillColor(org)
    c.rect(MX, y - 20, W - 2 * MX, 22, fill=1, stroke=0)
    c.setFillColor(HexColor('#FFFFFF'))
    c.setFont('Helvetica-Bold', 10)
    for i, h in enumerate(['Model', 'Input', 'Output', 'Monthly', 'Per day']):
        c.drawString(MX + 8 + i * colw, y - 14, h)
    y -= 20
    for ri, row in enumerate([('Claude Haiku 4.5', '$50.00', '$50.00', '$100.00', '$3.33'),
                              ('Claude Sonnet 4.6', '$150.00', '$150.00', '$300.00', '$10.00'),
                              ('Claude Opus 4.8', '$250.00', '$250.00', '$500.00', '$16.67')]):
        c.setFillColor(band if ri % 2 == 0 else HexColor('#FFFFFF'))
        c.rect(MX, y - 22, W - 2 * MX, 22, fill=1, stroke=0)
        for i, cell in enumerate(row):
            c.setFillColor(ink if i in (0, 3) else grey)
            c.setFont('Helvetica-Bold' if i in (0, 3) else 'Helvetica', 10)
            c.drawString(MX + 8 + i * colw, y - 15, cell)
        y -= 22
    c.setFillColor(grey)
    c.setFont('Helvetica', 9)
    c.drawString(MX, 50, 'etsy.com/shop/FranksMarketDesigns   ·   Unofficial and independent.')
    c.showPage()

    sheet('Rates — edit when pricing changes',
          'The only numbers you ever need to update. The Calculator reads from here.',
          ['Model', 'Input $/M', 'Output $/M', 'Cache $/M'],
          [(m[0], '$%.2f' % m[2], '$%.2f' % m[3], '$%.2f' % m[4]) for m in MODELS],
          input_rows=(0, 1, 2))
    c.setFillColor(grey)
    c.setFont('Helvetica', 9)
    c.drawString(MX, 50, 'Works in Excel, Numbers and Google Sheets.')
    c.showPage()
    c.save()
    print('Saved: %s' % path)


def main():
    os.makedirs(OUTDIR, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    build_calculator(wb)
    build_rates(wb)
    build_readme(wb)
    wb.active = 0
    wb.save(OUT)
    print('Saved: %s  (%d sheets)' % (OUT, len(wb.sheetnames)))
    print('   sheets: %s' % ', '.join(wb.sheetnames))
    build_preview_pdf()


if __name__ == '__main__':
    main()
